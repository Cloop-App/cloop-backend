"""
Add an import-audit sheet to the knowledge-graph workbook.

Writes a new file rather than editing in place: openpyxl rewrites a workbook it
opens, and silently drops anything it does not model (charts, images, some
formatting). The source workbook is left exactly as it was.
"""

import sys, json, collections, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="23386D")
WARN = Font(color="B8402A", bold=True)
OK = Font(color="2B6E4F", bold=True)


def write_sheet(wb, title, rows, widths):
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    for r in rows:
        ws.append(r)
    for c in range(1, len(rows[0]) + 1):
        ws.cell(row=1, column=c).font = HEAD
        ws.cell(row=1, column=c).fill = FILL
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                if cell.value.startswith("GAP") or cell.value.startswith("NOT "):
                    cell.font = WARN
                elif cell.value.startswith("OK"):
                    cell.font = OK
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def main(xlsx_in, json_path, xlsx_out):
    data = json.load(open(json_path, encoding="utf-8"))
    subjects, chapters, topics, concepts = (
        data["subjects"], data["chapters"], data["topics"], data["concepts"])

    wb = openpyxl.load_workbook(xlsx_in)

    # ── 1. what the importer actually loaded ─────────────────────────────
    v_ch = sum(1 for c in chapters if c["verified"])
    v_tp = sum(1 for t in topics if t["verified"])
    rows = [["Layer", "Rows imported", "Source-reconciled", "Unverified", "Status"],
            ["global_subjects", len(subjects), "-", "-", "OK loaded"],
            ["global_chapters", len(chapters), v_ch, len(chapters) - v_ch, "OK loaded"],
            ["global_topics", len(topics), v_tp, len(topics) - v_tp,
             "GAP 72% derived, not source-verified"],
            ["academic_concepts", len(concepts), 0, len(concepts),
             "GAP names only, no analytical layer"],
            ["concept_prerequisites", 0, 0, 0, "GAP column empty in Concept_Graph"],
            ["misconceptions", 0, 0, 0, "GAP column empty in Concept_Graph"]]
    write_sheet(wb, "V11_Import_Audit", rows, [26, 15, 19, 13, 44])

    # ── 2. coverage grid as imported ─────────────────────────────────────
    grid = collections.defaultdict(lambda: [0, 0, 0, 0, 0])
    by_key = {s["key"]: s for s in subjects}
    ch_by_key = {}
    for c in chapters:
        s = by_key[c["subject_key"]]
        g = grid[(s["board"], s["grade"])]
        g[1] += 1
        g[2] += 1 if c["verified"] else 0
        ch_by_key[c["chapter_key"]] = (s["board"], s["grade"])
    for s in subjects:
        grid[(s["board"], s["grade"])][0] += 1
    for t in topics:
        bg = ch_by_key.get(t["chapter_key"])
        if bg:
            grid[bg][3] += 1
            grid[bg][4] += 1 if t["verified"] else 0

    rows = [["Board / Exam", "Class", "Subjects", "Chapters", "Chapters verified",
             "Topics", "Topics verified", "Status"]]
    for (b, g), v in sorted(grid.items()):
        status = "OK sourced" if v[2] else "NOT a board syllabus (school-level)"
        if v[4] == 0 and v[2]:
            status = "OK chapters sourced; topics all derived"
        rows.append([b, g, v[0], v[1], v[2], v[3], v[4], status])
    for cls in ("Class 6", "Class 7"):
        rows.append(["Central Board of Secondary Education", cls, 0, 0, 0, 0, 0,
                     "GAP no rows in workbook"])
        rows.append(["Indian Certificate of Secondary Education", cls, 0, 0, 0, 0, 0,
                     "GAP no rows in workbook"])
    write_sheet(wb, "V11_Coverage_As_Imported", rows, [44, 13, 10, 10, 18, 9, 16, 42])

    # ── 3. the open gaps, stated plainly ─────────────────────────────────
    rows = [["#", "Gap", "Evidence", "Consequence", "What closes it"],
            [1, "Classes 6 and 7 absent",
             "Workbook covers VIII-XII only; Curriculum_Completeness lists 13 "
             "system/class rows, none below VIII",
             "Students in Classes 6-7 have no sourced curriculum; the live catalog "
             "serves them LLM-generated chapters",
             "Source CBSE/NCERT and CISCE syllabi for VI-VII"],
            [2, "ICSE Classes 8, 9 and 10 are identical",
             "All three classes share one chapter-list fingerprint "
             "(a945497e...), 43 chapters each",
             "A Class 9 student and a Class 10 student are taught the same syllabus; "
             "at most one can be right",
             "Separate the ICSE syllabus per class against the CISCE document"],
            [3, "No prerequisite graph",
             "Prerequisite Concept IDs filled on 0 of 3,151 rows; "
             "Prerequisite Concepts reads 'Not yet mapped - validation required'",
             "Mastery Engine v8 prerequisite gating cannot run",
             "Map prerequisites between concept codes"],
            [4, "No misconceptions",
             "Common Misconception filled on 0 of 3,151 rows",
             "Misconception detection and targeted remediation cannot run",
             "Author misconceptions per concept"],
            [5, "No difficulty calibration",
             "Difficulty reads 'To calibrate from item data' on all 3,151 rows",
             "Difficulty bands and adaptive item selection cannot run",
             "Calibrate from real item response data"],
            [6, "72% of topics unverified",
             "Topic_Master_v9: 3,161 of 4,390 marked "
             "'Derived from concept layer - requires curriculum validation'",
             "Only 1,229 topics may be presented as a board syllabus",
             "Validate derived topics against source documents"],
            [7, "CBSE Classes 9-10 have no verified topics",
             "44 and 41 chapters source-reconciled, but 0 topics verified in either",
             "Chapter structure is sourced; everything beneath it is derived",
             "Validate topics for the grades most of the pilot is in"],
            [8, "Class VIII is not a board syllabus",
             "Workbook's own note: 'Class VIII CBSE/ICSE - Scope-qualified - "
             "NOT A BOARD SYLLABUS'; 231 rows SCOPE-QUALIFIED",
             "Class 8 content is school-level, not examinable board content",
             "Keep the scope qualification visible in the product"],
            [9, "Workbook is not production-locked",
             "Curriculum_Data_Audit: 'Production readiness - Not yet locked'",
             "The workbook does not claim production readiness; "
             "neither should anything built on it",
             "Complete academic review, then lock a version"]]
    write_sheet(wb, "V11_Open_Gaps", rows, [5, 38, 54, 54, 44])

    # ── 4. provenance of this audit ──────────────────────────────────────
    rows = [["Item", "Value"],
            ["Audit produced by", "Cloop backend importer, tools/curriculum/"],
            ["Source workbook", xlsx_in.split("/")[-1]],
            ["Extraction", "tools/curriculum/extract.py -> data/curriculum/knowledge-graph-v10.json"],
            ["Import", "tools/curriculum/import.js (idempotent, non-destructive)"],
            ["Verified against", "PostgreSQL 16, production schema + 5 migrations"],
            ["Official sources checked", "NOT CHECKED - cbseacademic.nic.in, cisce.org, "
             "ncert.nic.in and jeemain.nta.nic.in are blocked by network policy"],
            ["Fact-check status", "NOT PERFORMED against board documents. Counts and "
             "internal consistency verified; syllabus accuracy is unverified."]]
    write_sheet(wb, "V11_Audit_Provenance", rows, [28, 96])

    wb.save(xlsx_out)
    print(f"written: {xlsx_out}")
    print("sheets added: V11_Import_Audit, V11_Coverage_As_Imported, "
          "V11_Open_Gaps, V11_Audit_Provenance")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
